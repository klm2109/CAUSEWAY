"""Build editable PowerPoint, browser slides and speaker notes from one slide source."""
from __future__ import annotations

import html
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]


def build() -> None:
    slides = json.loads((ROOT / "docs/paper2agent_slides.json").read_text(encoding="utf-8"))
    out = ROOT / "docs/meeting"
    out.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    notes = ["# Paper2Agent and the first TACO agent", "", "Approximately 10 minutes. Paper summary followed by our CAUSEWAY proposal.", ""]
    pages = []
    for number, item in enumerate(slides, 1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("142746")
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.14), presentation.slide_height)
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string("72DDBE")
        bar.line.fill.background()

        def text_box(x: float, y: float, w: float, h: float, text: str, size: int, color: str, bold: bool = False):
            frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
            frame.word_wrap = True
            for i, line in enumerate(text.split("\n")):
                paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(size)
                paragraph.font.bold = bold
                paragraph.font.color.rgb = RGBColor.from_string(color)
                paragraph.space_after = Pt(20)
            return frame

        text_box(.6, .35, 11.9, .4, "CAUSEWAY  /  LONG COVID ATLAS", 13, "72DDBE", True)
        text_box(.6, .95, 12.0, 1.25, item["title"], 34, "FFFFFF", True)
        text_box(.6, 2.08, 12.0, .6, item["subtitle"], 18, "92D9CF")
        text_box(.7, 2.95, 11.9, 3.65, "\n".join("• " + x for x in item["bullets"]), 21 if len(item["bullets"]) > 4 else 24, "E4ECF8")
        text_box(.6, 6.82, 11.5, .5, item["source"], 10, "B2C2D9")
        text_box(12.1, 6.8, .7, .4, f"{number:02}", 15, "72DDBE", True)
        slide.notes_slide.notes_text_frame.text = item["notes"] + "\n\nSource: " + item["source"]
        notes += [f"## Slide {number}: {item['title']}", "", *["- " + x for x in item["bullets"]], "", "Speaker notes: " + item["notes"], "", "Source: " + item["source"], ""]
        pages.append(f'<section class="slide" id="slide-{number}"><div class="brand">CAUSEWAY / LONG COVID ATLAS</div><h1>{html.escape(item["title"])}</h1><p class="subtitle">{html.escape(item["subtitle"])}</p><ul>{"".join("<li>" + html.escape(x) + "</li>" for x in item["bullets"])}</ul><details><summary>Speaker notes</summary><p>{html.escape(item["notes"])}</p></details><footer>{html.escape(item["source"])}<b>{number:02}</b></footer></section>')
    presentation.save(out / "Paper2Agent_CAUSEWAY.pptx")
    (out / "speaker_notes.md").write_text("\n".join(notes), encoding="utf-8")
    browser = '''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Paper2Agent × CAUSEWAY</title><style>
    *{box-sizing:border-box}body{margin:0;background:#0c182c;color:#e4ecf8;font:20px system-ui,sans-serif}.slide{position:relative;margin:28px auto;padding:40px 55px 72px;max-width:1200px;min-height:675px;background:#142746;border-left:10px solid #72ddbe;border-radius:8px}.brand{color:#72ddbe;font-size:13px;letter-spacing:.15em;font-weight:700}h1{font-size:40px;line-height:1.2;margin:35px 0 12px}.subtitle{color:#92d9cf;font-size:21px}ul{padding-left:28px;margin-top:36px}li{margin:20px 0;line-height:1.4}details{font-size:15px;color:#b2c2d9;margin:28px 0}summary{cursor:pointer}footer{position:absolute;left:55px;right:35px;bottom:22px;font-size:11px;color:#b2c2d9;overflow-wrap:anywhere}footer b{float:right;color:#72ddbe;font-size:16px}@media(max-width:700px){.slide{padding:25px 22px 80px;margin:12px;min-height:600px}h1{font-size:29px}body{font-size:17px}footer{left:22px}}@media print{@page{size:landscape;margin:0}.slide{break-after:page;margin:0;max-width:none;min-height:100vh;border-radius:0}details{display:none}}
    </style><body>''' + "\n".join(pages) + "</body></html>"
    (out / "Paper2Agent_CAUSEWAY.html").write_text(browser, encoding="utf-8")
    print(f"Built {len(slides)} slides, editable PowerPoint, browser deck and speaker notes in {out}")


if __name__ == "__main__":
    build()
